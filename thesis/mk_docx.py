"""Render Chapter3_revision.md as a Word document the thesis can absorb.

Times New Roman 12 pt, double-spaced, so drafted sections paste straight in.
Blockquotes are a presentation wrapper in the source, not content, so they
are stripped before the table and heading rules run -- otherwise every table
written as quoted draft text renders as a column of indented lines.
"""
import re
import sys

from docx import Document
from docx.shared import Inches, Pt

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

NL = chr(10)
src = open("Chapter3_revision.md", encoding="utf-8").read()
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 2.0

QUOTE = re.compile(r"^>\s?")
BOLD = re.compile(r"\*\*(.+?)\*\*")


def para(text, *, bold=False, size=12, before=0, after=0, spacing=2.0,
         indent=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = spacing
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    for i, chunk in enumerate(BOLD.split(text)):
        if not chunk:
            continue
        r = p.add_run(chunk.replace("`", ""))
        r.bold = bold or (i % 2 == 1)
        r.font.size = Pt(size)
        r.font.name = "Times New Roman"
    return p


def table(rows):
    t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = t.cell(ri, ci)
            c.text = ""
            pp = c.paragraphs[0]
            pp.paragraph_format.line_spacing = 1.0
            pp.paragraph_format.space_after = Pt(0)
            rr = pp.add_run(re.sub(r"[*`]", "", cell))
            rr.font.size = Pt(9.5)
            rr.font.name = "Times New Roman"
            rr.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


lines = src.split(NL)
i = 0
while i < len(lines):
    raw = lines[i].rstrip()
    quoted = raw.startswith(">")
    ln = QUOTE.sub("", raw) if quoted else raw
    ind = 0.35 if quoted else 0.0

    if ln.startswith("```"):
        i += 1
        block = []
        while i < len(lines) and not QUOTE.sub("", lines[i]).startswith("```"):
            block.append(QUOTE.sub("", lines[i].rstrip()))
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(NL.join(block))
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        i += 1
        continue

    if ln.startswith("|"):
        rows = []
        while i < len(lines):
            cur = QUOTE.sub("", lines[i].rstrip())
            if not cur.startswith("|"):
                break
            cells = [c.strip() for c in cur.strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            table(rows)
        continue

    if ln.startswith("# "):
        para(ln[2:], bold=True, size=14, before=12, after=8, spacing=1.0)
    elif ln.startswith("## "):
        para(ln[3:], bold=True, size=13, before=14, after=6, spacing=1.0)
    elif ln.startswith("### "):
        para(ln[4:], bold=True, size=12, before=10, after=4, spacing=1.0,
             indent=ind)
    elif ln.startswith(("- ", "* ")):
        para("• " + ln[2:], indent=ind + 0.3, spacing=1.5)
    elif ln.startswith("---"):
        para("_" * 60, spacing=1.0, after=6)
    elif ln.strip():
        para(ln, indent=ind)
    else:
        doc.add_paragraph().paragraph_format.line_spacing = 1.0
    i += 1

doc.save("Chapter3_revision.docx")
print("paragraphs:", len(doc.paragraphs), "| tables:", len(doc.tables))
